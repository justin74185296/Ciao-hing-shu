#!/usr/bin/env python3
"""
小紅書童裝供應商爬蟲 - 單文件版本
直接運行: python xhs_scraper.py

使用前請先：
1. pip install playwright playwright-stealth requests openpyxl pandas
2. playwright install chromium
3. 準備好小紅書 Cookie
"""

import asyncio
import json
import re
import os
import random
import time
import logging
from pathlib import Path
from datetime import datetime
from dataclasses import dataclass, field
from typing import Optional
from enum import Enum

# ========== 配置 ==========
COOKIE = ""  # 在這裡貼上你的 Cookie，或運行時輸入

KEYWORDS = [
    "童裝代理",
    "童裝一件代發",
    "韓版童裝批發招商",
]

MAX_PAGES = 5  # 每個關鍵字最大頁數
MAX_SUPPLIERS = 50  # 最大供應商數量

# ========== 日誌 ==========
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(message)s')
logger = logging.getLogger()

# ========== 數據模型 ==========
class SupplierType(Enum):
    UNKNOWN = "未知"
    FACTORY = "工廠/廠家"
    BRAND = "品牌方"
    DISTRIBUTOR = "經銷商/代理"

class ContactType(Enum):
    WECHAT = "微信"
    PHONE = "電話"

@dataclass
class ContactInfo:
    type: ContactType
    value: str
    source: str = ""

@dataclass
class NoteStats:
    likes: int = 0
    comments: int = 0
    collects: int = 0
    
    @property
    def total(self) -> int:
        return self.likes + self.comments * 3 + self.collects * 2

@dataclass
class Note:
    note_id: str
    title: str
    content: str = ""
    author_id: str = ""
    author_name: str = ""
    stats: NoteStats = field(default_factory=NoteStats)
    has_supplier_kw: bool = False

@dataclass
class Author:
    user_id: str
    nickname: str
    bio: str = ""
    followers: int = 0
    notes_count: int = 0
    ip_location: str = ""
    
    @property
    def url(self) -> str:
        return f"https://www.xiaohongshu.com/user/profile/{self.user_id}"

@dataclass
class Supplier:
    author: Author
    type: SupplierType = SupplierType.UNKNOWN
    contacts: list = field(default_factory=list)
    notes: list = field(default_factory=list)
    score: float = 0.0
    keywords: list = field(default_factory=list)
    
    def calc_score(self):
        score = 0
        # 互動分 (30)
        if self.notes:
            avg = sum(n.stats.total for n in self.notes) / len(self.notes)
            score += min(30, avg / 100 * 30)
        # 聯絡方式 (25)
        if any(c.type == ContactType.WECHAT for c in self.contacts):
            score += 25
        elif self.contacts:
            score += 15
        # 類型 (15)
        type_scores = {SupplierType.FACTORY: 15, SupplierType.BRAND: 12, SupplierType.DISTRIBUTOR: 8}
        score += type_scores.get(self.type, 0)
        # 活躍度 (15)
        if self.author.notes_count >= 50: score += 15
        elif self.author.notes_count >= 20: score += 10
        elif self.author.notes_count >= 10: score += 5
        # 筆記質量 (15)
        if self.notes:
            kw_notes = sum(1 for n in self.notes if n.has_supplier_kw)
            score += min(15, kw_notes / len(self.notes) * 15)
        self.score = round(score, 1)
        return self.score
    
    @property
    def wechat(self) -> str:
        for c in self.contacts:
            if c.type == ContactType.WECHAT:
                return c.value
        return ""
    
    def to_row(self) -> dict:
        return {
            "排名": 0,
            "暱稱": self.author.nickname,
            "品質評分": self.score,
            "供應商類型": self.type.value,
            "微信": self.wechat or "未找到",
            "粉絲數": self.author.followers,
            "筆記數": self.author.notes_count,
            "主頁連結": self.author.url,
            "簡介": self.author.bio[:80] if self.author.bio else "",
            "IP屬地": self.author.ip_location,
            "來源關鍵字": ", ".join(self.keywords),
        }

# ========== 工具函數 ==========
SUPPLIER_KW = ["代理", "招商", "一件代發", "批發", "工廠直供", "源頭", "廠家", "合作", "拿貨", "檔口", "一手貨源"]
FACTORY_KW = ["工廠", "廠家", "生產", "源頭"]
BRAND_KW = ["品牌", "原創", "設計師"]
EXCLUDE_KW = ["求推薦", "哪裡買", "測評", "開箱"]

WX_PATTERNS = [
    r'(?:微信|wx|vx|v信|威信|薇)[：:\s]*([a-zA-Z0-9_\-]{5,20})',
    r'[vV][：:\s]*([a-zA-Z0-9_]{6,20})',
]

def extract_wechat(text: str) -> list[ContactInfo]:
    contacts = []
    for pat in WX_PATTERNS:
        for m in re.finditer(pat, text, re.IGNORECASE):
            wxid = m.group(1).strip()
            if 5 <= len(wxid) <= 20 and any(c.isalpha() for c in wxid):
                if not any(x in wxid.lower() for x in ['xiaohongshu', 'weixin', 'wechat']):
                    contacts.append(ContactInfo(ContactType.WECHAT, wxid))
    seen = set()
    return [c for c in contacts if c.value.lower() not in seen and not seen.add(c.value.lower())]

def classify_supplier(bio: str, notes: list[Note]) -> SupplierType:
    text = bio + " " + " ".join(n.title + n.content for n in notes)
    text = text.lower()
    factory = sum(text.count(k) for k in FACTORY_KW)
    brand = sum(text.count(k) for k in BRAND_KW)
    if factory >= 2 and factory > brand: return SupplierType.FACTORY
    if brand >= 2: return SupplierType.BRAND
    if factory > 0 or brand > 0: return SupplierType.DISTRIBUTOR
    return SupplierType.UNKNOWN

def is_good_note(note: Note) -> bool:
    if note.stats.likes < 30: return False
    text = (note.title + note.content).lower()
    if any(k in text for k in EXCLUDE_KW): return False
    if any(k in text for k in SUPPLIER_KW):
        note.has_supplier_kw = True
        return True
    return False

# ========== 爬蟲 ==========
async def delay(a=1, b=3):
    await asyncio.sleep(random.uniform(a, b))

class Scraper:
    def __init__(self, cookie: str):
        self.cookie = cookie
        self.browser = None
        self.page = None
        self.suppliers: list[Supplier] = []
        self.seen_authors = set()
    
    async def init(self):
        from playwright.async_api import async_playwright
        from playwright_stealth import stealth_async
        
        logger.info("啟動瀏覽器...")
        pw = await async_playwright().start()
        self.browser = await pw.chromium.launch(headless=True)
        ctx = await self.browser.new_context(
            viewport={'width': 1920, 'height': 1080},
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 Chrome/120.0.0.0 Safari/537.36",
            locale='zh-CN'
        )
        # 設置 Cookie
        if self.cookie:
            cookies = []
            for item in self.cookie.split(';'):
                if '=' in item:
                    k, v = item.strip().split('=', 1)
                    cookies.append({'name': k, 'value': v, 'domain': '.xiaohongshu.com', 'path': '/'})
            await ctx.add_cookies(cookies)
        self.page = await ctx.new_page()
        await stealth_async(self.page)
        logger.info("瀏覽器就緒")
    
    async def close(self):
        if self.browser:
            await self.browser.close()
    
    async def search(self, keyword: str, max_pages: int):
        from urllib.parse import quote
        url = f"https://www.xiaohongshu.com/search_result?keyword={quote(keyword)}&source=web_search_result_notes"
        await self.page.goto(url, wait_until='networkidle')
        await delay(2, 4)
        
        for pg in range(1, max_pages + 1):
            logger.info(f"[{keyword}] 第 {pg}/{max_pages} 頁")
            try:
                await self.page.wait_for_selector('section.note-item, a.cover', timeout=10000)
            except:
                logger.warning("頁面加載超時")
                break
            
            # 提取數據
            notes = await self._get_notes()
            for note in notes:
                if not is_good_note(note): continue
                if note.author_id in self.seen_authors: continue
                
                # 獲取作者信息
                author = await self._get_author(note.author_id)
                if not author: continue
                
                self.seen_authors.add(note.author_id)
                
                # 構建供應商
                contacts = extract_wechat(author.bio + " " + note.title + " " + note.content)
                stype = classify_supplier(author.bio, [note])
                
                supplier = Supplier(
                    author=author, type=stype, contacts=contacts,
                    notes=[note], keywords=[keyword]
                )
                supplier.calc_score()
                
                if contacts or supplier.score > 30:
                    self.suppliers.append(supplier)
                    logger.info(f"✓ 發現: {author.nickname} ({supplier.score}分) {supplier.wechat or ''}")
                
                if len(self.suppliers) >= MAX_SUPPLIERS:
                    return
                
                await delay(1, 2)
            
            # 翻頁
            if pg < max_pages:
                await self.page.evaluate("window.scrollTo(0, document.body.scrollHeight)")
                await delay(3, 5)
    
    async def _get_notes(self) -> list[Note]:
        notes = []
        try:
            data = await self.page.evaluate("() => window.__INITIAL_STATE__")
            if data:
                items = data.get('search', {}).get('notes', {}).get('items', [])
                if not items:
                    items = data.get('searchResult', {}).get('items', [])
                for item in items:
                    nc = item.get('noteCard', item)
                    notes.append(Note(
                        note_id=item.get('id', ''),
                        title=nc.get('displayTitle', '') or nc.get('title', ''),
                        content=nc.get('desc', ''),
                        author_id=nc.get('user', {}).get('userId', ''),
                        author_name=nc.get('user', {}).get('nickname', ''),
                        stats=NoteStats(
                            likes=int(nc.get('interactInfo', {}).get('likedCount', 0) or 0),
                            comments=int(nc.get('interactInfo', {}).get('commentCount', 0) or 0),
                            collects=int(nc.get('interactInfo', {}).get('collectedCount', 0) or 0),
                        )
                    ))
        except Exception as e:
            logger.debug(f"提取失敗: {e}")
        return notes
    
    async def _get_author(self, uid: str) -> Optional[Author]:
        try:
            await self.page.goto(f"https://www.xiaohongshu.com/user/profile/{uid}", wait_until='networkidle')
            await delay(1.5, 3)
            data = await self.page.evaluate("() => window.__INITIAL_STATE__?.user")
            if data:
                info = data.get('userPageData', {})
                basic = info.get('basicInfo', {})
                inters = info.get('interactions', [])
                fans, notes = 0, 0
                for i in inters:
                    if i.get('type') == 'fans': fans = int(i.get('count', 0))
                    if i.get('type') == 'notes': notes = int(i.get('count', 0))
                return Author(
                    user_id=uid,
                    nickname=basic.get('nickname', ''),
                    bio=basic.get('desc', ''),
                    followers=fans,
                    notes_count=notes,
                    ip_location=basic.get('ipLocation', '')
                )
        except Exception as e:
            logger.error(f"獲取作者失敗: {e}")
        return None
    
    async def run(self, keywords: list[str], max_pages: int):
        await self.init()
        try:
            await self.page.goto("https://www.xiaohongshu.com", wait_until='networkidle')
            await delay(2, 3)
            
            for kw in keywords:
                logger.info(f"\n{'='*40}\n搜索: {kw}\n{'='*40}")
                await self.search(kw, max_pages)
                if len(self.suppliers) >= MAX_SUPPLIERS:
                    break
        finally:
            await self.close()
        return self.suppliers

# ========== 導出 ==========
def export_excel(suppliers: list[Supplier], filename: str = None):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        logger.error("請安裝 openpyxl: pip install openpyxl")
        return None
    
    if not filename:
        filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    
    suppliers = sorted(suppliers, key=lambda s: s.score, reverse=True)
    
    wb = Workbook()
    ws = wb.active
    ws.title = "供應商列表"
    
    headers = ["排名", "暱稱", "品質評分", "供應商類型", "微信", "粉絲數", "筆記數", "主頁連結", "簡介", "IP屬地", "來源關鍵字"]
    widths = [6, 16, 10, 12, 20, 10, 8, 45, 35, 10, 18]
    
    # 表頭
    for col, (h, w) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(1, col, h)
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal='center')
        ws.column_dimensions[get_column_letter(col)].width = w
    
    ws.freeze_panes = 'A2'
    
    # 數據
    for i, s in enumerate(suppliers, 1):
        row = s.to_row()
        row["排名"] = i
        for col, h in enumerate(headers, 1):
            val = row.get(h, "")
            cell = ws.cell(i + 1, col, val)
            if "連結" in h and val:
                cell.hyperlink = val
                cell.font = Font(color="0563C1", underline="single")
    
    os.makedirs("output", exist_ok=True)
    path = f"output/{filename}"
    wb.save(path)
    return path

def export_json(suppliers: list[Supplier], filename: str = None):
    if not filename:
        filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    
    suppliers = sorted(suppliers, key=lambda s: s.score, reverse=True)
    data = {
        "total": len(suppliers),
        "suppliers": [
            {
                "rank": i,
                "nickname": s.author.nickname,
                "score": s.score,
                "type": s.type.value,
                "wechat": s.wechat,
                "followers": s.author.followers,
                "url": s.author.url,
                "bio": s.author.bio,
            }
            for i, s in enumerate(suppliers, 1)
        ]
    }
    os.makedirs("output", exist_ok=True)
    path = f"output/{filename}"
    with open(path, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    return path

# ========== 主程序 ==========
def main():
    print("""
╔═══════════════════════════════════════════════════════╗
║     小紅書童裝供應商爬蟲 v1.0                          ║
╚═══════════════════════════════════════════════════════╝
    """)
    
    global COOKIE, KEYWORDS, MAX_PAGES, MAX_SUPPLIERS
    
    # 獲取 Cookie
    if not COOKIE:
        print("📌 請輸入小紅書 Cookie（直接貼上後按 Enter）：")
        print("   獲取方法：登錄 xiaohongshu.com → F12 → Network → 複製 Cookie")
        print()
        COOKIE = input("Cookie: ").strip()
    
    if not COOKIE:
        print("\n⚠️ 未輸入 Cookie，程序退出")
        return
    
    print(f"\n搜索關鍵字: {KEYWORDS}")
    print(f"最大頁數: {MAX_PAGES}")
    print(f"最大供應商數: {MAX_SUPPLIERS}")
    
    try:
        input("\n按 Enter 開始爬取...")
    except KeyboardInterrupt:
        print("\n已取消")
        return
    
    print("\n開始爬取...\n")
    
    # 運行爬蟲
    scraper = Scraper(COOKIE)
    try:
        suppliers = asyncio.run(scraper.run(KEYWORDS, MAX_PAGES))
    except KeyboardInterrupt:
        print("\n用戶中斷")
        suppliers = scraper.suppliers
    except Exception as e:
        print(f"\n錯誤: {e}")
        suppliers = scraper.suppliers
    
    if not suppliers:
        print("\n未找到供應商，請檢查 Cookie 是否有效")
        return
    
    # 導出
    print(f"\n✅ 找到 {len(suppliers)} 個供應商")
    
    excel_path = export_excel(suppliers)
    if excel_path:
        print(f"📊 Excel: {excel_path}")
    
    json_path = export_json(suppliers)
    print(f"📄 JSON: {json_path}")
    
    # 顯示 Top 5
    print("\n🏆 Top 5 供應商:")
    for i, s in enumerate(sorted(suppliers, key=lambda x: x.score, reverse=True)[:5], 1):
        print(f"   {i}. {s.author.nickname} ({s.score}分) - {s.wechat or '無微信'}")
    
    print(f"\n完成！結果保存在 output/ 資料夾")

if __name__ == "__main__":
    main()

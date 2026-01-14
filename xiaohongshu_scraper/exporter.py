"""
數據導出模組

功能：
- 導出為 Excel 格式
- 按品質評分排序
- 格式美化
- 多 Sheet 分類導出
"""

from pathlib import Path
from datetime import datetime
from typing import Optional
import json

try:
    import openpyxl
    from openpyxl import Workbook
    from openpyxl.styles import (
        Font, Fill, PatternFill, Alignment, Border, Side
    )
    from openpyxl.utils import get_column_letter
    from openpyxl.chart import BarChart, Reference
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    # 創建佔位類，避免類定義時出錯
    class PatternFill:
        def __init__(self, *args, **kwargs): pass
    class Font:
        def __init__(self, *args, **kwargs): pass
    class Alignment:
        def __init__(self, *args, **kwargs): pass
    class Workbook:
        def __init__(self, *args, **kwargs): pass
    def get_column_letter(x): return 'A'
    print("警告: openpyxl 未安裝，Excel 導出功能將不可用")

try:
    import pandas as pd
    PANDAS_AVAILABLE = True
except ImportError:
    PANDAS_AVAILABLE = False

from models import Supplier, SupplierType
from utils import logger


class ExcelExporter:
    """Excel 導出器"""
    
    # 列寬配置
    COLUMN_WIDTHS = {
        '排名': 8,
        '帳號ID': 20,
        '暱稱': 18,
        '主頁連結': 45,
        '品質評分': 12,
        '供應商類型': 14,
        '微信/聯絡方式': 25,
        '粉絲數': 12,
        '筆記數': 10,
        '簡介': 40,
        '最佳筆記連結': 45,
        '最佳筆記互動': 14,
        '來源關鍵字': 20,
        'IP屬地': 12,
        '爬取時間': 18,
    }
    
    # 顏色配置
    HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
    
    HIGH_SCORE_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
    MID_SCORE_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
    LOW_SCORE_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
    
    def __init__(self, output_dir: str = "output"):
        if not OPENPYXL_AVAILABLE:
            raise RuntimeError("openpyxl 未安裝，請運行: pip install openpyxl")
        
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
    
    def export(
        self, 
        suppliers: list[Supplier], 
        filename: str = None,
        sort_by: str = "quality_score",
        ascending: bool = False
    ) -> str:
        """
        導出供應商數據為 Excel
        
        Args:
            suppliers: 供應商列表
            filename: 文件名（不含擴展名）
            sort_by: 排序字段
            ascending: 是否升序
        
        Returns:
            輸出文件路徑
        """
        if not suppliers:
            logger.warning("沒有數據可導出")
            return ""
        
        # 生成文件名
        if not filename:
            filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        output_path = self.output_dir / f"{filename}.xlsx"
        
        # 排序
        suppliers = self._sort_suppliers(suppliers, sort_by, ascending)
        
        # 創建工作簿
        wb = Workbook()
        
        # 主 Sheet - 所有供應商
        self._create_main_sheet(wb, suppliers)
        
        # 分類 Sheet
        self._create_category_sheets(wb, suppliers)
        
        # 統計 Sheet
        self._create_stats_sheet(wb, suppliers)
        
        # 保存
        wb.save(output_path)
        logger.info(f"Excel 已導出: {output_path}")
        
        return str(output_path)
    
    def _sort_suppliers(
        self, 
        suppliers: list[Supplier], 
        sort_by: str,
        ascending: bool
    ) -> list[Supplier]:
        """排序供應商"""
        sort_keys = {
            'quality_score': lambda s: s.quality_score,
            'followers': lambda s: s.author.followers,
            'notes_count': lambda s: s.author.notes_count,
        }
        
        key_func = sort_keys.get(sort_by, sort_keys['quality_score'])
        return sorted(suppliers, key=key_func, reverse=not ascending)
    
    def _create_main_sheet(self, wb: Workbook, suppliers: list[Supplier]):
        """創建主 Sheet"""
        ws = wb.active
        ws.title = "全部供應商"
        
        # 表頭
        headers = list(self.COLUMN_WIDTHS.keys())
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal='center', vertical='center')
            
            # 設置列寬
            ws.column_dimensions[get_column_letter(col)].width = self.COLUMN_WIDTHS[header]
        
        # 凍結首行
        ws.freeze_panes = 'A2'
        
        # 數據行
        for row_idx, supplier in enumerate(suppliers, 2):
            row_data = supplier.to_excel_row()
            row_data['排名'] = row_idx - 1
            
            for col, header in enumerate(headers, 1):
                value = row_data.get(header, '')
                cell = ws.cell(row=row_idx, column=col, value=value)
                cell.alignment = Alignment(vertical='center', wrap_text=True)
                
                # 連結格式
                if '連結' in header and value:
                    cell.hyperlink = value
                    cell.font = Font(color="0563C1", underline="single")
            
            # 根據評分設置背景色
            score = supplier.quality_score
            if score >= 70:
                score_fill = self.HIGH_SCORE_FILL
            elif score >= 40:
                score_fill = self.MID_SCORE_FILL
            else:
                score_fill = self.LOW_SCORE_FILL
            
            ws.cell(row=row_idx, column=5).fill = score_fill
        
        # 設置自動篩選
        ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(suppliers) + 1}"
    
    def _create_category_sheets(self, wb: Workbook, suppliers: list[Supplier]):
        """按供應商類型創建分類 Sheet"""
        categories = {
            SupplierType.FACTORY: "工廠廠家",
            SupplierType.BRAND: "品牌方",
            SupplierType.DISTRIBUTOR: "經銷代理",
            SupplierType.INDIVIDUAL: "個人微商",
        }
        
        for supplier_type, sheet_name in categories.items():
            filtered = [s for s in suppliers if s.supplier_type == supplier_type]
            
            if not filtered:
                continue
            
            ws = wb.create_sheet(title=sheet_name)
            
            # 簡化的列
            simple_headers = ['排名', '暱稱', '品質評分', '微信/聯絡方式', '粉絲數', '主頁連結', '簡介']
            
            # 表頭
            for col, header in enumerate(simple_headers, 1):
                cell = ws.cell(row=1, column=col, value=header)
                cell.fill = self.HEADER_FILL
                cell.font = self.HEADER_FONT
                cell.alignment = Alignment(horizontal='center')
                ws.column_dimensions[get_column_letter(col)].width = self.COLUMN_WIDTHS.get(header, 15)
            
            # 數據
            for row_idx, supplier in enumerate(filtered, 2):
                row_data = supplier.to_excel_row()
                row_data['排名'] = row_idx - 1
                
                for col, header in enumerate(simple_headers, 1):
                    value = row_data.get(header, '')
                    cell = ws.cell(row=row_idx, column=col, value=value)
                    
                    if '連結' in header and value:
                        cell.hyperlink = value
                        cell.font = Font(color="0563C1", underline="single")
    
    def _create_stats_sheet(self, wb: Workbook, suppliers: list[Supplier]):
        """創建統計 Sheet"""
        ws = wb.create_sheet(title="統計分析")
        
        # 基本統計
        ws['A1'] = "統計項目"
        ws['B1'] = "數值"
        ws['A1'].font = Font(bold=True)
        ws['B1'].font = Font(bold=True)
        
        stats = [
            ("總供應商數", len(suppliers)),
            ("平均評分", round(sum(s.quality_score for s in suppliers) / len(suppliers), 1) if suppliers else 0),
            ("有微信聯絡方式", sum(1 for s in suppliers if s.primary_contact and '微信' in s.primary_contact)),
            ("工廠/廠家", sum(1 for s in suppliers if s.supplier_type == SupplierType.FACTORY)),
            ("品牌方", sum(1 for s in suppliers if s.supplier_type == SupplierType.BRAND)),
            ("經銷/代理", sum(1 for s in suppliers if s.supplier_type == SupplierType.DISTRIBUTOR)),
            ("個人微商", sum(1 for s in suppliers if s.supplier_type == SupplierType.INDIVIDUAL)),
            ("高評分(70+)", sum(1 for s in suppliers if s.quality_score >= 70)),
            ("中評分(40-70)", sum(1 for s in suppliers if 40 <= s.quality_score < 70)),
            ("低評分(<40)", sum(1 for s in suppliers if s.quality_score < 40)),
        ]
        
        for row, (name, value) in enumerate(stats, 2):
            ws.cell(row=row, column=1, value=name)
            ws.cell(row=row, column=2, value=value)
        
        # 設置列寬
        ws.column_dimensions['A'].width = 20
        ws.column_dimensions['B'].width = 15
        
        # 添加圖表（如果數據足夠）
        if len(suppliers) >= 3:
            self._add_chart(ws, len(stats) + 3)
    
    def _add_chart(self, ws, start_row: int):
        """添加圖表"""
        # 類型分佈數據
        ws.cell(row=start_row, column=1, value="供應商類型分佈")
        ws.cell(row=start_row, column=1).font = Font(bold=True, size=12)
        
        chart = BarChart()
        chart.type = "col"
        chart.style = 10
        chart.title = "供應商類型分佈"
        chart.y_axis.title = "數量"
        
        # 數據範圍 (從統計數據中讀取)
        data = Reference(ws, min_col=2, min_row=5, max_row=8)
        cats = Reference(ws, min_col=1, min_row=5, max_row=8)
        
        chart.add_data(data, titles_from_data=False)
        chart.set_categories(cats)
        chart.shape = 4
        
        ws.add_chart(chart, f"D{start_row}")


class CSVExporter:
    """CSV 導出器（備用）"""
    
    def __init__(self, output_dir: str = "output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
    
    def export(self, suppliers: list[Supplier], filename: str = None) -> str:
        """導出為 CSV"""
        if not suppliers:
            return ""
        
        if not filename:
            filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        output_path = self.output_dir / f"{filename}.csv"
        
        # 排序
        suppliers = sorted(suppliers, key=lambda s: s.quality_score, reverse=True)
        
        # 轉換為行數據
        rows = []
        for idx, supplier in enumerate(suppliers, 1):
            row = supplier.to_excel_row()
            row['排名'] = idx
            rows.append(row)
        
        # 使用 pandas 或手動寫入
        if PANDAS_AVAILABLE:
            df = pd.DataFrame(rows)
            df.to_csv(output_path, index=False, encoding='utf-8-sig')
        else:
            import csv
            with open(output_path, 'w', encoding='utf-8-sig', newline='') as f:
                if rows:
                    writer = csv.DictWriter(f, fieldnames=rows[0].keys())
                    writer.writeheader()
                    writer.writerows(rows)
        
        logger.info(f"CSV 已導出: {output_path}")
        return str(output_path)


class JSONExporter:
    """JSON 導出器"""
    
    def __init__(self, output_dir: str = "output"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(exist_ok=True)
    
    def export(self, suppliers: list[Supplier], filename: str = None) -> str:
        """導出為 JSON"""
        if not suppliers:
            return ""
        
        if not filename:
            filename = f"童裝供應商_{datetime.now().strftime('%Y%m%d_%H%M%S')}"
        
        output_path = self.output_dir / f"{filename}.json"
        
        # 排序
        suppliers = sorted(suppliers, key=lambda s: s.quality_score, reverse=True)
        
        # 轉換為字典
        data = {
            "export_time": datetime.now().isoformat(),
            "total_count": len(suppliers),
            "suppliers": [s.to_dict() for s in suppliers]
        }
        
        with open(output_path, 'w', encoding='utf-8') as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        
        logger.info(f"JSON 已導出: {output_path}")
        return str(output_path)


def export_suppliers(
    suppliers: list[Supplier],
    format: str = "excel",
    output_dir: str = "output",
    filename: str = None
) -> str:
    """
    導出供應商數據
    
    Args:
        suppliers: 供應商列表
        format: 導出格式 (excel/csv/json)
        output_dir: 輸出目錄
        filename: 文件名
    
    Returns:
        輸出文件路徑
    """
    exporters = {
        'excel': ExcelExporter,
        'csv': CSVExporter,
        'json': JSONExporter,
    }
    
    exporter_class = exporters.get(format.lower())
    if not exporter_class:
        raise ValueError(f"不支援的格式: {format}")
    
    try:
        exporter = exporter_class(output_dir)
        return exporter.export(suppliers, filename)
    except Exception as e:
        logger.error(f"導出失敗: {e}")
        
        # 備用方案
        if format == 'excel' and not OPENPYXL_AVAILABLE:
            logger.info("嘗試使用 CSV 格式")
            return CSVExporter(output_dir).export(suppliers, filename)
        
        raise


if __name__ == "__main__":
    # 測試導出
    from models import Author, Note, NoteStats, ContactInfo, ContactType
    
    # 創建測試數據
    test_suppliers = []
    
    for i in range(5):
        author = Author(
            user_id=f"user_{i}",
            nickname=f"童裝供應商{i+1}",
            bio=f"專業童裝批發，微信：factory{i+1}",
            followers=1000 * (i + 1),
            notes_count=50 + i * 10,
            ip_location="廣東"
        )
        
        note = Note(
            note_id=f"note_{i}",
            title=f"童裝批發招商{i+1}",
            stats=NoteStats(likes=100*(i+1), comments=20*(i+1), collects=50*(i+1))
        )
        note.has_supplier_keywords = True
        
        supplier = Supplier(
            author=author,
            supplier_type=SupplierType.FACTORY if i % 2 == 0 else SupplierType.DISTRIBUTOR,
            contacts=[ContactInfo(ContactType.WECHAT, f"factory{i+1}", "簡介")],
            notes=[note],
            source_keywords=["童裝代理", "童裝批發"]
        )
        supplier.calculate_quality_score()
        test_suppliers.append(supplier)
    
    # 測試 Excel 導出
    print("=== 測試 Excel 導出 ===")
    try:
        path = export_suppliers(test_suppliers, format='excel', filename='test_export')
        print(f"Excel 導出成功: {path}")
    except Exception as e:
        print(f"Excel 導出失敗: {e}")
    
    # 測試 JSON 導出
    print("\n=== 測試 JSON 導出 ===")
    path = export_suppliers(test_suppliers, format='json', filename='test_export')
    print(f"JSON 導出成功: {path}")
    
    # 測試 CSV 導出
    print("\n=== 測試 CSV 導出 ===")
    path = export_suppliers(test_suppliers, format='csv', filename='test_export')
    print(f"CSV 導出成功: {path}")
